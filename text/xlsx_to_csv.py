import sys
import os
# https://openpyxl.readthedocs.io/en/stable/tutorial.html#installation
from openpyxl import load_workbook
import csv


# migrate translated text from xlsx to csv
TARGETS = [
    {"xlsx" : "M_00_09.xlsx",
     "csv" : "M_00_09.BIN.csv",
     "changes" : [
         (25,2,458),
         (488,461,35),
         (528,497,29),
         (561,527,3),
         (566,530,16),
         (586,547,15),
         (603,563,66),
         (673,630,19),
         (693,650,16),
         (713,667,6),
         (722,674,3),
         (729,677,253),
         (998,931,177),
         (1177,1109,81),
         (1283,1191,25),
         (1317,1217,9),
         (1336,1227,15),
         (1366,1243,246),
         (2,1490,22),
         (484,1513,4),
         (558,1517,3),
         (564,1520,2),
         (583,1522,3),
         (602,1525,1),
         (710,1526,3),
         (720,1529,2),
         (725,1531,4),
         (983,1535,14),
         (1259,1550,23),
         (1309,1574,8),
         (1327,1582,8),
         (1352,1591,13)
         ]},
    {"xlsx" : "M_01_01.xlsx",
     "csv" : "M_01_01.BIN.csv",
     "changes" : [
         (2,2,3448)
         ]},
    {"xlsx" : "M_02_01.xlsx",
     "csv" : "M_02_01.BIN.csv",
     "changes" : [
         ()
        ]},
    {"xlsx" : "",
     "csv" : "M_02_02.BIN.csv",
     "changes" : [
         ()
        ]},
    {"xlsx" : "",
     "csv" : "M_03_01.BIN.csv",
     "changes" : [
         ()
        ]},
    {"xlsx" : "",
     "csv" : "M_04_01.BIN.csv",
     "changes" : [
         ()
        ]},
    {"xlsx" : "M_CMN.xlsx",
     "csv" : "M_CMN.BIN.csv",
     "changes" : [
         (2,2,5455),
         (5533,5458,3),
         (5458,5461,65),
         (5525,5527,3),
         (5529,5530,2)
        ]},
    {"xlsx" : "SLPS_25574.xlsx",
     "csv" : "SLPS_255.74.csv",
     "changes" : [
         (2,2,438),
         (444,441,1683),
         (441,2281,3),
         (2128,2284,11),
         (2198,2296,1),
         (2199,2298,84),
         (2284,2384,10)
        ]}
    ]


for target in TARGETS:
    csv_path = target["csv"]
    xlsx_path = target["xlsx"]
    try:
        wb = load_workbook(filename=xlsx_path)
        print("Writing to " + csv_path)
    except:
        continue
    ws = wb.active
    csv_data = []
    with open(csv_path, 'r', encoding='utf-8') as csvfile:
        reader = csv.DictReader(csvfile, delimiter = ';', quotechar = '"')
        for row in reader:
            csv_data.append([row["block"],
                             row["j_offset"],
                             row["j_string"],
                             row["e_string"],
                             row["pointer"]])
    changes = target["changes"]
    for item in changes:
        cell = item[0]
        line = item[1] - 2
        length = item[2] + 1
        if target["xlsx"] =="M_00_09.xlsx":
            for i in range(length):
                csv_data[line + i][3] = ws["D{}".format(cell + i)].value
        else:
            for i in range(length):
                csv_data[line + i][3] = ws["C{}".format(cell + i)].value
    with open(csv_path, 'w', encoding='utf-8', newline='') as csvfile:
        writer = csv.writer(csvfile, delimiter=';', quotechar='"', quoting=csv.QUOTE_ALL)
        writer.writerow(["block",
                         "j_offset",
                         "j_string",
                         "e_string",
                         "pointer"])
        for line in csv_data:
            writer.writerow(line)
input("Press ENTER to close.")
